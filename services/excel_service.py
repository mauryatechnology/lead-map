import io
import os
from typing import List
from models.schemas import Lead
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter


COLUMNS = ["Name", "City", "Instagram", "Phone", "Website", "Notes", "About"]
FIELD_MAP = ["name", "city", "instagram", "phone", "website", "notes", "about"]

HEADER_COLOR = "1A1A2E"
ACCENT_COLOR = "E94560"
ROW_ALT = "F8F9FF"
ROW_NORMAL = "FFFFFF"
BORDER_COLOR = "DEE2F0"


def generate_excel(leads: List[Lead]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"

    # Title row
    ws.merge_cells("A1:G1")
    title_cell = ws["A1"]
    title_cell.value = "📍 Google Maps Business Leads"
    title_cell.font = Font(bold=True, size=14, color="FFFFFF", name="Arial")
    title_cell.fill = PatternFill("solid", fgColor=HEADER_COLOR)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    # Subtitle
    ws.merge_cells("A2:G2")
    sub_cell = ws["A2"]
    from datetime import datetime
    sub_cell.value = f"Extracted: {datetime.now().strftime('%Y-%m-%d %H:%M')}  |  Total Leads: {len(leads)}"
    sub_cell.font = Font(italic=True, size=10, color="888888", name="Arial")
    sub_cell.fill = PatternFill("solid", fgColor="F0F1F8")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # Header row
    header_row = 3
    thin = Side(style="thin", color=BORDER_COLOR)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx, col_name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.value = col_name
        cell.font = Font(bold=True, size=11, color="FFFFFF", name="Arial")
        cell.fill = PatternFill("solid", fgColor=ACCENT_COLOR)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[header_row].height = 28

    # Data rows
    seen = set()
    row_num = header_row + 1
    for lead in leads:
        key = f"{lead.name}|{lead.phone}|{lead.website}"
        if key in seen:
            continue
        seen.add(key)

        fill_color = ROW_ALT if (row_num - header_row) % 2 == 0 else ROW_NORMAL
        row_fill = PatternFill("solid", fgColor=fill_color)

        for col_idx, field in enumerate(FIELD_MAP, start=1):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.value = getattr(lead, field, "")
            cell.font = Font(size=10, name="Arial")
            cell.fill = row_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = border

        ws.row_dimensions[row_num].height = 20
        row_num += 1

    # Column widths
    col_widths = [32, 18, 22, 18, 35, 20, 45]
    for i, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # Freeze header
    ws.freeze_panes = f"A{header_row + 1}"

    # Auto filter
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(COLUMNS))}{row_num - 1}"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def save_excel(leads: List[Lead], path: str = "gym_leads.xlsx"):
    data = generate_excel(leads)
    with open(path, "wb") as f:
        f.write(data)
    return path
